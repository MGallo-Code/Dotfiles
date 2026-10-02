"""Programmatic DOCX generation from a typed spec via python-docx."""

from __future__ import annotations

from pathlib import Path
from typing import Annotated, Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from mcp.server.fastmcp import FastMCP
from pydantic import Field

from docgen.models import (
    Block,
    DocxSpec,
    HeadingBlock,
    ImageBlock,
    InlineStyle,
    ListBlock,
    PageBreakBlock,
    ParagraphBlock,
    TableBlock,
)
from docgen.paths import PathError, resolve_input_path, resolve_output_path
from docgen.utils import err, ok

_ALIGN_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


def _apply_inline_style(run, style: InlineStyle | None) -> None:
    if style is None:
        return
    if style.bold is not None:
        run.bold = style.bold
    if style.italic is not None:
        run.italic = style.italic
    if style.size_pt is not None:
        run.font.size = Pt(style.size_pt)


def _add_heading(doc, block: HeadingBlock) -> None:
    if block.style_name:
        para = doc.add_paragraph(block.text, style=block.style_name)
    else:
        para = doc.add_heading(block.text, level=block.level)
    if not block.style_name and block.level == 0:
        return
    # No-op; heading is already added.


def _add_paragraph(doc, block: ParagraphBlock) -> None:
    para = doc.add_paragraph(style=block.style_name) if block.style_name else doc.add_paragraph()
    run = para.add_run(block.text)
    if block.style and block.style.align in _ALIGN_MAP:
        para.alignment = _ALIGN_MAP[block.style.align]
    _apply_inline_style(run, block.style)


def _add_list(doc, block: ListBlock) -> None:
    style = "List Number" if block.ordered else "List Bullet"
    for item in block.items:
        doc.add_paragraph(item, style=style)


def _draw_table_borders(table) -> None:
    """Apply single-line borders to every edge via raw OXML.

    Fallback for when the document has no "Table Grid" style (common in
    third-party / exported reference docs): keeps tables visually gridded
    without depending on a named style being present.
    """
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tbl_pr.append(borders)


def _apply_grid_style(table) -> None:
    """Style a table as a grid, tolerating reference docs that lack the style.

    `table.style = "Table Grid"` raises KeyError when the (possibly
    reference-seeded) document defines no style by that name. Fall back to
    manually drawn borders so the call never crashes the whole render.
    """
    try:
        table.style = "Table Grid"
    except KeyError:
        _draw_table_borders(table)


def _add_table(doc, block: TableBlock) -> None:
    if not block.rows:
        return
    cols = max(len(r) for r in block.rows)
    table = doc.add_table(rows=len(block.rows), cols=cols)
    _apply_grid_style(table)
    for r_idx, row in enumerate(block.rows):
        for c_idx in range(cols):
            cell_text = row[c_idx] if c_idx < len(row) else ""
            cell = table.cell(r_idx, c_idx)
            cell.text = cell_text
            if block.header and r_idx == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True


def _add_image(doc, block: ImageBlock) -> None:
    path = resolve_input_path(block.path)
    if block.width_in is not None:
        doc.add_picture(str(path), width=Inches(block.width_in))
    else:
        doc.add_picture(str(path))


def _add_page_break(doc, _: PageBreakBlock) -> None:
    run = doc.add_paragraph().add_run()
    run.add_break(WD_BREAK.PAGE)


_DISPATCH = {
    "heading": _add_heading,
    "paragraph": _add_paragraph,
    "list": _add_list,
    "table": _add_table,
    "image": _add_image,
    "page_break": _add_page_break,
}


def _build_doc(spec: DocxSpec, reference_doc: Path | None):
    doc = Document(str(reference_doc)) if reference_doc else Document()
    for block in spec.blocks:
        handler = _DISPATCH[block.type]
        handler(doc, block)
    return doc


def register_docx_build_tools(server: FastMCP) -> None:
    @server.tool(
        name="docx_create_from_spec",
        description=(
            "Build a .docx programmatically from a structured spec. "
            "Spec has shape {blocks: [...]}. Block types: heading, paragraph, list, "
            "table, image, page_break. Optionally pass reference_doc to seed styles "
            "from an existing .docx (e.g. brand template). "
            "Table cells are strings only in v0; multi-paragraph cells not yet supported."
        ),
    )
    async def docx_create_from_spec(
        spec: Annotated[dict[str, Any], Field(description='Spec dict, e.g. {"blocks": [{"type": "heading", "level": 1, "text": "Hello"}]}')],
        output_path: Annotated[str, Field(description="Destination .docx path")],
        reference_doc: Annotated[str | None, Field(description="Optional path to a reference .docx whose styles seed the output")] = None,
        overwrite: bool = True,
    ) -> dict[str, Any]:
        try:
            parsed = DocxSpec.model_validate(spec)
        except Exception as e:  # noqa: BLE001 — surface validation errors cleanly
            return err(f"Invalid spec: {e}")

        try:
            target = resolve_output_path(output_path, overwrite=overwrite)
            ref = resolve_input_path(reference_doc) if reference_doc else None
        except PathError as e:
            return err(str(e))

        try:
            doc = _build_doc(parsed, ref)
            doc.save(str(target))
        except PathError as e:
            return err(str(e))
        except Exception as e:  # noqa: BLE001
            return err(f"{type(e).__name__}: {e}")

        return ok(path=str(target), bytes_written=target.stat().st_size, blocks=len(parsed.blocks))
