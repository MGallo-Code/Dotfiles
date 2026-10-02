import os
import tempfile

# Tests write to a throwaway output dir, never the user's real one. Set before docgen.paths
# is imported (conftest loads first).
os.environ.setdefault("DOCGEN_OUTPUT_DIR", tempfile.mkdtemp(prefix="docgen-test-"))

import zipfile
from pathlib import Path

import pytest
from docx import Document


def pytest_addoption(parser):
    parser.addoption(
        "--slow",
        action="store_true",
        default=False,
        help="Run slow PDF/Playwright tests",
    )


def pytest_collection_modifyitems(config, items):
    if config.getoption("--slow"):
        return
    skip_slow = pytest.mark.skip(reason="need --slow option to run")
    for item in items:
        if "slow" in item.keywords:
            item.add_marker(skip_slow)


@pytest.fixture
def docx_template_with_name(tmp_path: Path) -> Path:
    """Build a tiny .docx template with a Jinja2 {{name}} placeholder."""
    doc = Document()
    doc.add_heading("Hello {{ name }}", level=1)
    doc.add_paragraph("Greetings, {{ name }}!")
    out = tmp_path / "tpl_name.docx"
    doc.save(out)
    return out


@pytest.fixture
def docx_xml():
    """Return a helper that extracts word/document.xml from a .docx path."""

    def _extract(path: Path) -> str:
        with zipfile.ZipFile(path) as zf:
            return zf.read("word/document.xml").decode("utf-8")

    return _extract
