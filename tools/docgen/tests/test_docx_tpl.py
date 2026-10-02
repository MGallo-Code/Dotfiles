from __future__ import annotations

import asyncio
from pathlib import Path

import pytest

from docgen.docx_tpl import register_docx_tpl_tools
from docgen.server import mcp


@pytest.fixture(scope="module", autouse=True)
def register_tools():
    register_docx_tpl_tools(mcp)


def _registered(name: str):
    return mcp._tool_manager.get_tool(name).fn


def call_tool(name: str, **kwargs):
    return asyncio.run(_registered(name)(**kwargs))


def test_inspect_template_variables(docx_template_with_name: Path):
    res = call_tool("inspect_docx_template", template_path=str(docx_template_with_name))
    assert res["ok"], res
    assert "name" in res["undeclared_variables"]


def test_fill_template_writes_substituted_text(docx_template_with_name: Path, tmp_path: Path, docx_xml):
    out = tmp_path / "filled.docx"
    res = call_tool(
        "docx_from_template",
        template_path=str(docx_template_with_name),
        context={"name": "Michael"},
        output_path=str(out),
    )
    assert res["ok"], res
    xml = docx_xml(Path(res["path"]))
    assert "Michael" in xml
    assert "{{ name }}" not in xml


def test_fill_template_inline_image(docx_template_with_name: Path, tmp_path: Path):
    """docxtpl InlineImage support: pass {path, width_mm} dict in context."""
    from PIL import Image

    img_path = tmp_path / "logo.png"
    Image.new("RGB", (20, 20), "blue").save(img_path)

    # Build a template that has an {{ img }} placeholder we can fill with an image.
    from docx import Document

    tpl_doc = Document()
    tpl_doc.add_paragraph("Header")
    tpl_doc.add_paragraph("{{ img }}")  # placeholder gets replaced by InlineImage
    tpl_path = tmp_path / "tpl_img.docx"
    tpl_doc.save(tpl_path)

    out = tmp_path / "with_image.docx"
    res = call_tool(
        "docx_from_template",
        template_path=str(tpl_path),
        context={"img": {"path": str(img_path), "width_mm": 30}},
        output_path=str(out),
    )
    assert res["ok"], res

    import zipfile

    with zipfile.ZipFile(res["path"]) as zf:
        media = [n for n in zf.namelist() if n.startswith("word/media/")]
    assert media, "no embedded image"
